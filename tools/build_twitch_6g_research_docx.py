#!/usr/bin/env python3
"""Build the verified Twitch/iPod 6G failure-analysis report."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs/twitch-6g-deep-research/Twitch_6G_Failure_Analysis.docx"
PURPLE = "9146FF"
PURPLE_DARK = "6441A5"
CHARCOAL = "18181B"
GRAY = "53535F"
LIGHT = "F4F4F6"
WHITE = "FFFFFF"
GREEN = "157347"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    fill = tc_pr.find(qn("w:shd"))
    if fill is None:
        fill = OxmlElement("w:shd")
        tc_pr.append(fill)
    fill.set(qn("w:fill"), color)


def margins(section):
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("ROCKPOD  •  TWITCH / IPOD 6G   |   ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def configure(doc):
    section = doc.sections[0]
    margins(section)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.25)
    add_page_number(section.footer.paragraphs[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string(CHARCOAL)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (
        ("Title", 28, CHARCOAL),
        ("Subtitle", 11, GRAY),
        ("Heading 1", 17, PURPLE_DARK),
        ("Heading 2", 12, CHARCOAL),
        ("Heading 3", 10, PURPLE_DARK),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)


def set_cell_text(cell, text, *, bold=False, color=CHARCOAL, size=8.2):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1.5)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    tbl.autofit = False
    for i, header in enumerate(headers):
        set_cell_text(tbl.rows[0].cells[i], header, bold=True, color=WHITE)
        shade(tbl.rows[0].cells[i], PURPLE_DARK)
        if widths:
            tbl.rows[0].cells[i].width = Inches(widths[i])
    for row_index, values in enumerate(rows):
        cells = tbl.add_row().cells
        for i, value in enumerate(values):
            set_cell_text(cells[i], value)
            if row_index % 2:
                shade(cells[i], LIGHT)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return tbl


def label(doc, title, value, color=PURPLE):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    left, right = tbl.rows[0].cells
    set_cell_text(left, title.upper(), bold=True, color=WHITE, size=8)
    set_cell_text(right, value, bold=True, color=CHARCOAL, size=9)
    shade(left, color)
    shade(right, LIGHT)
    left.width = Inches(1.6)
    right.width = Inches(5.7)
    return tbl


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(2)
        number = p.add_run(f"{index}.  ")
        number.bold = True
        p.add_run(item)


def page(doc):
    doc.add_page_break()


def build():
    doc = Document()
    configure(doc)
    props = doc.core_properties
    props.title = "Twitch on iPod 6G: Failure Analysis"
    props.subject = "H.264/AAC, RockPod sync, and composite-output qualification"
    props.author = "RockPod engineering"
    props.keywords = "Twitch, iPod 6G, Rockbox, H.264, AAC, composite"

    accent = doc.add_table(rows=1, cols=1)
    accent.style = "Table Grid"
    accent.rows[0].height = Inches(0.12)
    shade(accent.cell(0, 0), PURPLE)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.add_run("TWITCH  /  IPOD 6G").bold = True
    p.runs[0].font.size = Pt(10)
    p.runs[0].font.color.rgb = RGBColor.from_string(PURPLE)
    title = doc.add_paragraph(style="Title")
    title.add_run("Playback, audio, sync, and composite failure analysis")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run(
        "Evidence-backed implementation report  •  5 September 2026  •  ipod6g only"
    )
    doc.add_paragraph()
    label(doc, "Primary finding", "Good AAC audio was corrupted by a signed 2 GiB file-offset boundary.")
    label(doc, "Playback policy", "H.264 first; converted AAC-LC audio; MPEG only as recovery.", GREEN)
    label(doc, "Composite finding", "One scanned YUV buffer was being rewritten; use field-edge double buffering.")
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.add_run("Reserved validation VOD\n").bold = True
    r = p.add_run("https://www.twitch.tv/videos/2860855364")
    r.font.color.rgb = RGBColor.from_string(PURPLE_DARK)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.add_run(
        "Decision status: host conversion and shared AAC decoder qualify. "
        "Physical VPU, analog composite, and full audio-lifecycle gates remain observable-only on the deployed 6G."
    )

    page(doc)
    doc.add_heading("Executive finding", level=1)
    doc.add_paragraph(
        "The Emma H.264 files did not lose their audio during download or host conversion. "
        "The synced files contain AAC-LC 44.1 kHz stereo, and Rockbox's own AAC codec produces non-silent PCM from them."
    )
    doc.add_heading("Two device defects explained the report", level=2)
    numbered(doc, [
        "The AAC chunk table stored valid unsigned 32-bit MP4 positions but returned them through a signed int. At or above byte 2,147,483,648, later audio positions became negative and interleaved AAC navigation failed.",
        "Composite output rewrote the one planar YUV buffer currently being scanned by the TV processor, allowing an analog field to contain rows from two framebuffer states.",
    ])
    doc.add_heading("Implemented decision", level=2)
    bullets(doc, [
        "Keep Apple-contract H.264 as the default Twitch format and first launch choice.",
        "Decode and re-encode Twitch source audio to AAC-LC, 44,100 Hz, stereo; never silently drop a source audio track.",
        "Carry AAC file positions as signed 64-bit values while preserving Rockbox's supported unsigned 32-bit MP4 range.",
        "Cache an MPEG sibling as a last-resort per-VOD recovery path.",
        "Continue sync after creator, download, H.264, MPEG, thumbnail, or chat failure; preserve an older playable copy.",
        "Publish a complete inactive composite frame and switch its three planes at a bounded field edge.",
    ])
    doc.add_heading("Important content fact", level=2)
    doc.add_paragraph(
        "VOD 2860855364 is silent in the Twitch source for approximately its first 190.87 seconds. "
        "The full stream is intentionally preserved. A device test must seek beyond 600 seconds before treating silence as a decoder failure."
    )

    page(doc)
    doc.add_heading("Supplied VOD evidence", level=1)
    table(doc, ["Artifact", "Duration", "Size", "Streams"], [
        ["Downloaded source", "8,109.383 s", "1,389,143,197 B", "852×480 H.264 Main / AAC-LC 48 kHz stereo"],
        ["RockPod H.264", "8,109.355 s", "1,561,827,252 B", "Apple H.264 / AAC-LC 44.1 kHz stereo"],
        ["Clean MPEG recovery", "8,109.375 s", "552,396,800 B", "320×240 MPEG-2 30 fps / MP2 44.1 kHz stereo"],
    ], [1.55, 1.05, 1.35, 3.35])
    doc.add_heading("Recovery checksum", level=2)
    p = doc.add_paragraph()
    r = p.add_run("6bf68aa20792aecb988352b55765fb98c48e3b617e0f631b0e028be8d61bcb8a")
    r.font.name = "Aptos Mono"
    r.font.size = Pt(8)
    doc.add_heading("Emma inventory scope", level=2)
    doc.add_paragraph(
        "The local library contains six VODs under creator aliases emma and emmablackerytv. "
        "Four already had device H.264 copies and two older entries had only MPEG copies from earlier partial runs. "
        "The H.264-first policy now applies to all six. Unrelated sweet_anita rows are excluded from the Emma inventory, although the decoder correction also covers them."
    )
    doc.add_heading("Why host playback looked correct", level=2)
    doc.add_paragraph(
        "Desktop players use wide file offsets and follow every MP4 chunk correctly. Rockbox's table also retained the right unsigned values; only the helper's signed return type damaged values after 2 GiB. "
        "That combination explains why the same file could sound normal on a PC and fail on the iPod at a later live join or seek."
    )

    page(doc)
    doc.add_heading("Failure gap matrix", level=1)
    table(doc, ["Failure", "Evidence / cause", "Control", "Physical gate"], [
        ["Late or live-join audio missing", "AAC packet positions exceed signed 32-bit range", "64-bit signed helper/caller", "Late seek + 60 s listen"],
        ["Silent opening", "Source VOD has a 190.87 s silent intro", "Preserve full media; test later interval", "Seek beyond 600 s"],
        ["H.264 fails without recovery", "VPU/AAC failure was not actionable", "Explicit error + same-ID MPEG", "Force fallback once"],
        ["One item aborts whole sync", "Operation-wide exception boundary", "Per-creator and per-VOD isolation", "Failure injection"],
        ["Retries reconvert MPEG", "No persistent completed cache", "Atomic signature cache", "Second-sync cache hit"],
        ["Composite horizontal lines", "Scanned YUV buffer rewritten in place", "Field-edge double buffering", "30 s motion test"],
        ["Chat overlaps video", "Independent overlay boundary", "One animated shared boundary", "Select open/close"],
    ], [1.45, 2.15, 1.85, 1.85])
    doc.add_paragraph(
        "The MPEG fallback is intentionally subordinate: a successful H.264 launch never routes through MPEG. "
        "A backup failure is reported without invalidating the primary H.264 copy."
    )

    page(doc)
    doc.add_heading("Large-file AAC acceptance", level=1)
    doc.add_paragraph(
        "Each row below was tested at a known non-silent interval whose first observed AAC packet was beyond byte 2,147,483,647. "
        "The simulator exercises the same MP4 parser, libm4a lookup, AAC codec, PCM ring, and playback mixer used by the device."
    )
    table(doc, ["VOD", "File size", "Start", "AAC position", "Result"], [
        ["2857240220", "2,283,328,465", "11,800 s", "2,271,061,838", "10 s pass"],
        ["2858925101", "2,179,532,988", "11,100 s", "2,169,250,615", "10 s pass"],
        ["2862600186", "2,171,177,441", "11,100 s", "2,159,998,910", "10 s pass"],
        ["2864411857", "2,335,023,638", "12,000 s", "2,302,661,400", "30 s + 10 s pass"],
    ], [1.25, 1.55, 1.05, 1.55, 1.7])
    doc.add_heading("Automated status", level=2)
    label(doc, "Focused tests", "38 passed, 1 skipped", GREEN)
    label(doc, "6G hardware build", "Linked successfully", GREEN)
    label(doc, "Composite static gate", "Passed", GREEN)
    doc.add_paragraph()
    doc.add_paragraph(
        "The simulator does not execute the S5L8702 VPU or the analog output block. Its success qualifies the shared software audio path, not physical H.264 decoding or composite presentation."
    )

    page(doc)
    doc.add_heading("Implemented runtime contract", level=1)
    bullets(doc, [
        "Request the best complete Twitch video+audio rendition at or below 480p.",
        "Reject a download that lacks either a video or an audio stream.",
        "Prepare Apple-contract H.264 first and map source audio into a new AAC-LC 44.1 kHz stereo track zero.",
        "Reject an H.264 or MPEG output that dropped audio present in the source.",
        "Retain a stable-ID M4V plus an MPG recovery sibling; prefer M4V, MP4, MOV, then MPG at launch.",
        "Open the MPG sibling only after an explicit H.264 VPU/AAC failure.",
        "Treat every creator and VOD as its own failure boundary; later items continue and prior playable media survives.",
        "Use cached MPEG conversion so a retry does not repeat completed work.",
        "Keep chat in sidecars with real Twitch/Twemoji assets; the full-height rail reflows rather than overlays video.",
        "Keep composite buffers target-owned and independent from plugin, codec, shared audio, and playback memory.",
    ])
    doc.add_heading("Copyright-muted content", level=2)
    doc.add_paragraph(
        "Conversion can preserve and normalize audio that exists in the acquired Twitch rendition. It cannot reconstruct audio Twitch removed or replaced before delivery. No implementation claims recovery of absent copyrighted audio."
    )

    page(doc)
    doc.add_heading("Physical 6G acceptance checklist", level=1)
    numbered(doc, [
        "Reboot into the deployed firmware. Open Applications → Twitch and confirm there is no file/plugin-open error.",
        "Play VOD 2860855364. Its first ~191 seconds are source-silent; seek beyond 600 seconds and confirm audio, volume, pause/resume, and Menu exit.",
        "Play an Emma M4V larger than 2 GiB, seek beyond the table's position, and listen continuously for at least 60 seconds.",
        "Join an Emma live channel at a late wall-clock position and confirm AAC remains the playback clock.",
        "Press Select. Verify the 146 px full-height chat rail slides smoothly and continuously reduces video width.",
        "With composite enabled, scroll and toggle chat for 30 seconds. Reject mixed rows, horizontal tearing, stale frames, or field roll.",
        "Force one H.264 failure and verify only the same VOD ID's MPEG sibling opens.",
        "Exercise fresh boot → Twitch, Database music → Twitch, Files music → Twitch, and return to both music paths without mute or freeze.",
    ])
    doc.add_heading("Release boundary", level=2)
    doc.add_paragraph(
        "Do not label VPU, composite, or lifecycle behavior physically qualified until these observations pass on the actual 6G. Build success and simulator audio are necessary evidence, not a substitute for hardware."
    )

    page(doc)
    doc.add_heading("Sources and implementation trail", level=1)
    doc.add_heading("Authoritative external sources", level=2)
    table(doc, ["Source", "Use"], [
        ["Apple iPod classic 160 GB (Late 2009) technical specifications\nhttps://support.apple.com/en-ie/112601", "H.264 Baseline/Low Complexity and AAC-LC playback limits"],
        ["FFmpeg documentation\nhttps://ffmpeg.org/ffmpeg.html", "Optional stream-map semantics; why an audio map may legally disappear"],
        ["Linux kernel DRM/KMS documentation\nhttps://docs.kernel.org/gpu/drm-kms.html", "Authoritative corroboration for vblank/page-flip scanout synchronization"],
    ], [4.75, 2.55])
    doc.add_heading("Repository evidence", level=2)
    table(doc, ["Area", "Files"], [
        ["Audio lifecycle", "docs/plugin-audio-lifecycle-steering.md; apps/video_audio.c; apps/video_pcm.c"],
        ["MP4/AAC", "lib/rbcodec/codecs/aac.c; lib/rbcodec/codecs/libm4a/m4a.c"],
        ["H.264 player", "apps/video_playback.c; apps/plugins/openh264_player.c"],
        ["Twitch app/sync", "apps/plugins/twitch.c; rockpod/services/twitch_app.py; app_video_sync.py; video_rvp.py"],
        ["Composite", "firmware/target/arm/s5l8702/ipod6g/videoout-6g.c; docs/ipod6g-dcp750-videoout-results.md"],
        ["Regression", "tools/h264_audio_sim_gate.py; rockpod/tests/test_h264_audio_large_file_source.py"],
    ], [1.55, 5.75])
    doc.add_heading("Conclusion", level=2)
    doc.add_paragraph(
        "The missing-audio report is explained by a deterministic large-file integer boundary plus genuine silent source intros—not by absent converted tracks. "
        "The correction covers all compatible AAC/M4A files below Rockbox's 4 GiB limit, keeps H.264 primary, and leaves MPEG as recovery. Host and simulator evidence passes; the deployed 6G remains the authority for VPU and composite acceptance."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
